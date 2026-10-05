"""
MCP Connector — Local files & Windows filesystem.

Covers: PDF, Word (.docx), Excel (.xlsx), CSV, Markdown, PowerPoint (.pptx),
plain text, and generic folder / filesystem operations (list, create,
move, copy, delete) on local disk — including any Windows path
(C:\\Users\\... etc.) or network share the machine running this server
can see.

No external service credentials needed. Optionally set FILES_ROOT_DIR in
.env to give relative paths a default root; absolute paths always work
as given regardless of that setting.

Every underlying library here (pypdf, python-docx, openpyxl, python-pptx)
is synchronous, so each tool wraps its call in asyncio.to_thread — this
keeps the server non-blocking: a slow read of a large spreadsheet doesn't
stall other requests being handled concurrently.
"""

import argparse
import difflib
import asyncio
import csv as csv_module
import io
import os
import re
import shutil
import tempfile
from contextlib import contextmanager
from pathlib import Path

from env_config import load_project_env
from mcp.server.fastmcp import FastMCP

load_project_env(Path(__file__).resolve().parent / ".env")

mcp = FastMCP(
    name="files-connector",
    instructions=(
        "CRUD operations on local files (PDF, Word, Excel, CSV, Markdown, "
        "PowerPoint, plain text) and folders/filesystem paths."
    ),
)

FILES_ROOT_DIR = os.environ.get("FILES_ROOT_DIR", "").strip()


def _resolve(path: str) -> Path:
    p = Path(path)
    if not p.is_absolute() and FILES_ROOT_DIR:
        p = Path(FILES_ROOT_DIR) / p
    return p


def _missing_file_message(path: Path) -> str:
    if not path.parent.exists():
        return f"File not found: {path}"

    candidates = [entry.name for entry in path.parent.iterdir() if entry.is_file()]
    close_matches = difflib.get_close_matches(path.name, candidates, n=3, cutoff=0.6)
    if close_matches:
        suggestions = ", ".join(str(path.parent / name) for name in close_matches)
        return f"File not found: {path}. Did you mean: {suggestions}"

    return f"File not found: {path}"


@contextmanager
def _copied_source(path: Path):
    if not path.exists():
        raise FileNotFoundError(_missing_file_message(path))
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir) / path.name
        shutil.copy2(path, temp_path)
        yield temp_path


# --- Filesystem / folders (Local folders, Windows file system) -----------

@mcp.tool()
async def fs_list_dir(path: str) -> list:
    """List entries in a folder, with type (file/dir) and size."""
    def _run():
        p = _resolve(path)
        return [
            {"name": e.name, "type": "dir" if e.is_dir() else "file",
             "size": e.stat().st_size if e.is_file() else None}
            for e in sorted(p.iterdir())
        ]
    return await asyncio.to_thread(_run)


@mcp.tool()
async def fs_get_info(path: str) -> dict:
    """Get metadata (exists, type, size, modified time) for a path."""
    def _run():
        p = _resolve(path)
        if not p.exists():
            return {"exists": False, "path": str(p)}
        st = p.stat()
        return {
            "exists": True, "path": str(p), "type": "dir" if p.is_dir() else "file",
            "size": st.st_size, "modified": st.st_mtime,
        }
    return await asyncio.to_thread(_run)


@mcp.tool()
async def fs_create_dir(path: str) -> dict:
    """Create a folder (and any missing parent folders)."""
    def _run():
        p = _resolve(path)
        p.mkdir(parents=True, exist_ok=True)
        return {"created_dir": str(p)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def fs_move(src: str, dst: str) -> dict:
    """Move or rename a file/folder."""
    def _run():
        s, d = _resolve(src), _resolve(dst)
        shutil.move(str(s), str(d))
        return {"moved_to": str(d)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def fs_copy(src: str, dst: str) -> dict:
    """Copy a file (or a folder, recursively)."""
    def _run():
        s, d = _resolve(src), _resolve(dst)
        if s.is_dir():
            shutil.copytree(s, d, dirs_exist_ok=True)
        else:
            shutil.copy2(s, d)
        return {"copied_to": str(d)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def fs_delete(path: str, recursive: bool = False) -> dict:
    """Delete a file, or a folder (recursive=true required for non-empty folders). Irreversible."""
    def _run():
        p = _resolve(path)
        if p.is_dir():
            if recursive:
                shutil.rmtree(p)
            else:
                p.rmdir()
        else:
            p.unlink()
        return {"deleted": str(p)}
    return await asyncio.to_thread(_run)


# --- Plain text / Markdown ------------------------------------------------

@mcp.tool()
async def text_read(path: str) -> str:
    """Read a plain text or Markdown file's full content."""
    def _run():
        source = _resolve(path)
        with _copied_source(source) as source_copy:
            return source_copy.read_text(encoding="utf-8", errors="replace")
    return await asyncio.to_thread(_run)


@mcp.tool()
async def text_write(path: str, content: str, mode: str = "overwrite") -> dict:
    """Write to a plain text or Markdown file. mode: 'overwrite' (default, creates if missing) or 'append'."""
    def _run():
        p = _resolve(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        if mode == "append":
            with p.open("a", encoding="utf-8") as f:
                f.write(content)
        else:
            p.write_text(content, encoding="utf-8")
        return {"written": str(p), "mode": mode}
    return await asyncio.to_thread(_run)


# --- CSV -------------------------------------------------------------------

@mcp.tool()
async def csv_read(path: str, max_rows: int = 100) -> list:
    """Read a CSV file as a list of row dicts (header row used as keys)."""
    def _run():
        p = _resolve(path)
        with _copied_source(p) as source_copy:
            with source_copy.open(newline="", encoding="utf-8", errors="replace") as f:
                reader = csv_module.DictReader(f)
                rows = []
                for i, row in enumerate(reader):
                    if i >= max_rows:
                        break
                    rows.append(row)
                return rows
    return await asyncio.to_thread(_run)


@mcp.tool()
async def csv_write(path: str, rows: list[dict]) -> dict:
    """Create or overwrite a CSV file from a list of row dicts. Column order follows the first row's keys."""
    def _run():
        p = _resolve(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        if not rows:
            p.write_text("", encoding="utf-8")
            return {"written": str(p), "rows": 0}
        with p.open("w", newline="", encoding="utf-8") as f:
            writer = csv_module.DictWriter(f, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        return {"written": str(p), "rows": len(rows)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def csv_append_row(path: str, row: dict) -> dict:
    """Append one row to an existing CSV file."""
    def _run():
        p = _resolve(path)
        file_exists = p.exists() and p.stat().st_size > 0
        with p.open("a", newline="", encoding="utf-8") as f:
            writer = csv_module.DictWriter(f, fieldnames=list(row.keys()))
            if not file_exists:
                writer.writeheader()
            writer.writerow(row)
        return {"appended_to": str(p)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def csv_to_xlsx(path: str, output_path: str | None = None, sheet_name: str = "Sheet1") -> dict:
    """
    Convert a CSV file into a new Excel workbook. output_path defaults to
    the same name and folder as the CSV, with a .xlsx extension — pass it
    explicitly to write somewhere else instead.
    """
    def _run():
        from openpyxl import Workbook
        p = _resolve(path)
        with _copied_source(p) as source_copy:
            with source_copy.open(newline="", encoding="utf-8", errors="replace") as f:
                rows = list(csv_module.reader(f))
        dest = _resolve(output_path) if output_path else p.with_suffix(".xlsx")
        dest.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in rows:
            ws.append(row)
        wb.save(str(dest))
        return {"created": str(dest), "rows": len(rows)}
    return await asyncio.to_thread(_run)


# --- PDF ---------------------------------------------------------------


def _normalize_pdf_text(text: str) -> str:
    """Collapse pypdf's word-by-word layout into readable prose."""
    return re.sub(r"\s+", " ", text).strip()

@mcp.tool()
async def pdf_read_text(path: str, max_pages: int | None = None) -> str:
    """Extract text from a PDF (pages joined with form-feed markers). Won't work on scanned/image-only PDFs."""
    def _run():
        from pypdf import PdfReader
        source = _resolve(path)
        with _copied_source(source) as source_copy:
            reader = PdfReader(str(source_copy))
            pages = reader.pages[:max_pages] if max_pages else reader.pages
            return "\n\x0c\n".join(_normalize_pdf_text(p.extract_text() or "") for p in pages)
    return await asyncio.to_thread(_run)


@mcp.tool()
async def pdf_get_info(path: str) -> dict:
    """Get page count and basic metadata for a PDF."""
    def _run():
        from pypdf import PdfReader
        source = _resolve(path)
        with _copied_source(source) as source_copy:
            reader = PdfReader(str(source_copy))
            meta = reader.metadata or {}
            return {"page_count": len(reader.pages), "title": meta.get("/Title"), "author": meta.get("/Author")}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def pdf_create(path: str, pages: list[str]) -> dict:
    """Create a new PDF, one page of plain text per list entry."""
    def _run():
        from fpdf import FPDF
        p = _resolve(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        pdf = FPDF()
        pdf.set_font("Helvetica", size=12)
        for text in pages:
            pdf.add_page()
            pdf.multi_cell(0, 8, text)
        pdf.output(str(p))
        return {"created": str(p), "pages": len(pages)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def pdf_append_page(path: str, text: str) -> dict:
    """Append a new text page to an existing PDF."""
    def _run():
        from fpdf import FPDF
        from pypdf import PdfReader, PdfWriter
        p = _resolve(path)
        with _copied_source(p) as source_copy:
            new_page_pdf = FPDF()
            new_page_pdf.set_font("Helvetica", size=12)
            new_page_pdf.add_page()
            new_page_pdf.multi_cell(0, 8, text)
            new_bytes = io.BytesIO(new_page_pdf.output())

            writer = PdfWriter()
            for pg in PdfReader(str(source_copy)).pages:
                writer.add_page(pg)
            for pg in PdfReader(new_bytes).pages:
                writer.add_page(pg)
            with p.open("wb") as f:
                writer.write(f)
            return {"appended_to": str(p)}
    return await asyncio.to_thread(_run)


# --- Word (.docx) --------------------------------------------------------


def _docx_paragraph_text(paragraph) -> str:
    from docx.oxml.ns import qn

    parts: list[str] = []
    for node in paragraph._p.iter():
        if node.tag == qn("w:t"):
            if node.text:
                parts.append(node.text)
        elif node.tag == qn("w:tab"):
            parts.append("\t")
        elif node.tag in {qn("w:br"), qn("w:cr")}:
            parts.append("\n")
    return "".join(parts).strip()


def _docx_cell_text(cell) -> str:
    texts: list[str] = []
    for block in _iter_docx_blocks(cell):
        if block[0] == "paragraph":
            text = _docx_paragraph_text(block[1])
            if text:
                texts.append(text)
        elif block[0] == "table":
            for row in block[1]:
                row_text = " | ".join(cell_text for cell_text in row if cell_text)
                if row_text:
                    texts.append(row_text)
    return "\n".join(texts).strip()


def _docx_table_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        row_texts: list[str] = []
        for cell in row.cells:
            row_texts.append(_docx_cell_text(cell))
        rows.append(row_texts)
    return rows


def _iter_docx_blocks(parent):
    from docx.document import Document as DocumentType
    from docx.table import _Cell, Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    if isinstance(parent, DocumentType):
        parent_element = parent.element.body
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
    else:
        parent_element = parent

    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield ("paragraph", Paragraph(child, parent))
        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            yield ("table", _docx_table_rows(table))


def _iter_docx_paragraphs(parent):
    from docx.document import Document as DocumentType
    from docx.table import _Cell, Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    if isinstance(parent, DocumentType):
        parent_element = parent.element.body
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
    else:
        parent_element = parent

    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            for row in table.rows:
                for cell in row.cells:
                    yield from _iter_docx_paragraphs(cell)


def _docx_read_result(doc, source_path: Path, structured: bool) -> str | dict:
    blocks = []
    texts = []
    for block_type, value in _iter_docx_blocks(doc):
        if block_type == "paragraph":
            text = _docx_paragraph_text(value)
            if text:
                texts.append(text)
                if structured:
                    blocks.append({"type": "paragraph", "text": text})
        elif block_type == "table":
            flat_rows = []
            for row in value:
                row_values = [cell_text for cell_text in row if cell_text]
                if row_values:
                    texts.extend(row_values)
                flat_rows.append(row)
            if structured:
                blocks.append({"type": "table", "rows": flat_rows})

    flat_text = "\n".join(texts)
    if structured:
        return {
            "path": str(source_path),
            "blocks": blocks,
            "text": flat_text,
        }
    return flat_text

@mcp.tool()
async def docx_read(path: str, structured: bool = False) -> str | dict:
    """Read text from a Word document; structured=True preserves paragraph/table layout."""
    def _run():
        from docx import Document
        source_path = _resolve(path)
        with _copied_source(source_path) as source_copy:
            doc = Document(str(source_copy))
            return _docx_read_result(doc, source_path, structured)
    return await asyncio.to_thread(_run)


@mcp.tool()
async def docx_create(path: str, paragraphs: list[str], title: str | None = None) -> dict:
    """Create a new Word document from a list of paragraphs, with an optional heading."""
    def _run():
        from docx import Document
        p = _resolve(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        if title:
            doc.add_heading(title, level=1)
        for para in paragraphs:
            doc.add_paragraph(para)
        doc.save(str(p))
        return {"created": str(p), "paragraphs": len(paragraphs)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def docx_append_paragraph(path: str, text: str) -> dict:
    """Append a paragraph to an existing Word document."""
    def _run():
        from docx import Document
        p = _resolve(path)
        doc = Document(str(p))
        doc.add_paragraph(text)

        doc.save(str(p))
        return {"appended_to": str(p)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def docx_replace_text(path: str, find: str, replace: str) -> dict:
    """Find-and-replace text across paragraphs, tables, and text boxes in a Word document."""
    def _run():
        from docx.oxml.ns import qn
        from docx import Document
        p = _resolve(path)
        with _copied_source(p) as source_copy:
            doc = Document(str(source_copy))
            count = 0
            for para in _iter_docx_paragraphs(doc):
                for node in para._p.iter():
                    if node.tag == qn("w:t") and node.text and find in node.text:
                        node.text = node.text.replace(find, replace)
                        count += 1
            doc.save(str(p))
            return {"replacements_made": count}
    return await asyncio.to_thread(_run)


# --- Excel (.xlsx) -------------------------------------------------------

@mcp.tool()
async def xlsx_read(path: str, sheet: str | None = None, max_rows: int = 100) -> list:
    """Read rows from an Excel sheet (default: active sheet)."""
    def _run():
        from openpyxl import load_workbook
        source = _resolve(path)
        with _copied_source(source) as source_copy:
            wb = load_workbook(str(source_copy), read_only=True, data_only=True)
            try:
                ws = wb[sheet] if sheet else wb.active
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= max_rows:
                        break
                    rows.append(list(row))
                return rows
            finally:
                wb.close()
    return await asyncio.to_thread(_run)


@mcp.tool()
async def xlsx_create(path: str, rows: list[list], sheet_name: str = "Sheet1") -> dict:
    """Create a new Excel workbook from a list of rows (each row is a list of cell values)."""
    def _run():
        from openpyxl import Workbook
        p = _resolve(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in rows:
            ws.append(row)
        wb.save(str(p))
        return {"created": str(p), "rows": len(rows)}
    return await asyncio.to_thread(_run)


@mcp.tool()
async def xlsx_update_cell(path: str, cell: str, value, sheet: str | None = None) -> dict:
    """Update one cell's value, e.g. cell='B2'. Note: this drops formulas/formatting in re-saved cells."""
    def _run():
        from openpyxl import load_workbook
        p = _resolve(path)
        with _copied_source(p) as source_copy:
            wb = load_workbook(str(source_copy))
            try:
                ws = wb[sheet] if sheet else wb.active
                ws[cell] = value
                wb.save(str(p))
                return {"updated_cell": cell, "value": value}
            finally:
                wb.close()
    return await asyncio.to_thread(_run)


@mcp.tool()
async def xlsx_append_row(path: str, row: list, sheet: str | None = None) -> dict:
    """Append a row to an existing Excel sheet."""
    def _run():
        from openpyxl import load_workbook
        p = _resolve(path)
        with _copied_source(p) as source_copy:
            wb = load_workbook(str(source_copy))
            try:
                ws = wb[sheet] if sheet else wb.active
                ws.append(row)
                wb.save(str(p))
                return {"appended_to": str(p)}
            finally:
                wb.close()
    return await asyncio.to_thread(_run)


# --- PowerPoint (.pptx) ----------------------------------------------------


def _pptx_part_text_frames(shapes):
    for shape in shapes:
        if getattr(shape, "has_text_frame", False):
            yield shape.text_frame
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def _pptx_iter_parts(prs, include_notes: bool = True):
    for slide_number, slide in enumerate(prs.slides, start=1):
        yield {
            "kind": "slide",
            "slide_number": slide_number,
            "owner": slide,
            "shapes": slide.shapes,
        }
        if include_notes and slide.has_notes_slide:
            notes_slide = slide.notes_slide
            yield {
                "kind": "notes",
                "slide_number": slide_number,
                "owner": notes_slide,
                "shapes": notes_slide.shapes,
            }

    for master_index, master in enumerate(prs.slide_masters, start=1):
        yield {
            "kind": "master",
            "master_index": master_index,
            "owner": master,
            "shapes": master.shapes,
        }
        for layout_index, layout in enumerate(master.slide_layouts, start=1):
            yield {
                "kind": "layout",
                "master_index": master_index,
                "layout_index": layout_index,
                "owner": layout,
                "shapes": layout.shapes,
            }


def _pptx_collect_text(prs, structured: bool = False) -> list | dict:
    slides = []
    notes = []
    masters = []
    layouts = []

    for part in _pptx_iter_parts(prs):
        text = "\n".join(text_frame.text for text_frame in _pptx_part_text_frames(part["shapes"]))
        entry = {"text": text}

        if part["kind"] == "slide":
            entry["slide"] = part["slide_number"]
            slides.append(entry)
        elif part["kind"] == "notes":
            entry["slide"] = part["slide_number"]
            notes.append(entry)
        elif part["kind"] == "master":
            entry["master"] = part["master_index"]
            masters.append(entry)
        elif part["kind"] == "layout":
            entry["master"] = part["master_index"]
            entry["layout"] = part["layout_index"]
            layouts.append(entry)

    if structured:
        return {
            "slides": slides,
            "notes": notes,
            "masters": masters,
            "layouts": layouts,
        }
    return slides


def _pptx_default_box(prs, target: str, position: dict | None = None):
    from pptx.util import Inches

    emu_per_inch = 914400
    slide_width = prs.slide_width / emu_per_inch
    slide_height = prs.slide_height / emu_per_inch
    target = target.lower()

    if position is None:
        if target == "header":
            position = {"left": 0.5, "top": 0.2, "width": slide_width - 1.0, "height": 0.5}
        elif target == "footer":
            position = {"left": 0.5, "top": slide_height - 0.8, "width": slide_width - 1.0, "height": 0.5}
        elif target in {"master", "layout"}:
            position = {"left": 0.8, "top": 1.0, "width": slide_width - 1.6, "height": 0.8}
        elif target == "notes":
            position = {"left": 0.8, "top": 0.8, "width": slide_width - 1.6, "height": 0.8}
        else:
            position = {"left": 0.8, "top": 1.25, "width": slide_width - 1.6, "height": slide_height - 2.0}

    return (
        Inches(float(position.get("left", 0.0))),
        Inches(float(position.get("top", 0.0))),
        Inches(float(position.get("width", slide_width - 1.6))),
        Inches(float(position.get("height", slide_height - 2.0))),
    )


def _pptx_replace_in_paragraph(paragraph, pattern, replace_text: str) -> int:
    runs = paragraph.runs
    matches = list(pattern.finditer("".join(run.text for run in runs)))
    for match in reversed(matches):
        offset = 0
        first_run = last_run = None
        for index, run in enumerate(runs):
            end = offset + len(run.text)
            if first_run is None and match.start() < end:
                first_run = index
                start_offset = match.start() - offset
            if first_run is not None and match.end() <= end:
                last_run = index
                end_offset = match.end() - offset
                break
            offset = end
        if first_run is None or last_run is None:
            continue
        if first_run == last_run:
            run = runs[first_run]
            run.text = run.text[:start_offset] + replace_text + run.text[end_offset:]
        else:
            runs[first_run].text = runs[first_run].text[:start_offset] + replace_text
            for index in range(first_run + 1, last_run):
                runs[index].text = ""
            runs[last_run].text = runs[last_run].text[end_offset:]
    return len(matches)


def _pptx_apply_edits(prs, edits: list[dict]) -> int:
    from pptx.util import Pt

    def _selected_parts(edit: dict):
        target = str(edit.get("target", "slide")).lower()
        slide_number = edit.get("slide_number")
        master_index = edit.get("master_index")
        layout_index = edit.get("layout_index")

        for part in _pptx_iter_parts(prs):
            kind = part["kind"]
            if target == "all":
                yield part
                continue

            if target in {"slide", "slides", "header", "footer"}:
                if kind != "slide":
                    continue
                if slide_number is not None and part["slide_number"] != int(slide_number):
                    continue
            elif target in {"notes", "note", "speaker_notes"}:
                if kind != "notes":
                    continue
                if slide_number is not None and part["slide_number"] != int(slide_number):
                    continue
            elif target in {"master", "masters", "slide_master"}:
                if kind != "master":
                    continue
                if master_index is not None and part["master_index"] != int(master_index):
                    continue
            elif target in {"layout", "layouts", "slide_layout"}:
                if kind != "layout":
                    continue
                if master_index is not None and part["master_index"] != int(master_index):
                    continue
                if layout_index is not None and part["layout_index"] != int(layout_index):
                    continue
            else:
                raise ValueError(
                    "target must be one of: slide, notes, master, layout, header, footer, or all"
                )

            yield part

    applied = 0
    for edit in edits:
        action = str(edit.get("action", "edit")).lower()
        action = {
            "update": "edit",
            "replace": "edit",
            "insert": "add",
            "append": "add",
            "remove": "delete",
        }.get(action, action)

        if action not in {"add", "edit", "delete"}:
            raise ValueError("action must be add, edit, or delete")

        target = str(edit.get("target", "slide")).lower()
        parts = list(_selected_parts(edit))
        if not parts:
            continue

        if action == "add":
            text = edit.get("text")
            if text is None:
                raise ValueError("add operations require a text value")

            font_size = edit.get("font_size")
            bold = edit.get("bold")
            position = edit.get("position")
            for part in parts:
                left, top, width, height = _pptx_default_box(prs, target, position)
                shape_collection = part["owner"].shapes
                if hasattr(shape_collection, "add_textbox"):
                    shape = shape_collection.add_textbox(left, top, width, height)
                else:
                    shape_id = shape_collection._next_shape_id
                    shape_name = f"TextBox {shape_id - 1}"
                    textbox = shape_collection._spTree.add_textbox(
                        shape_id,
                        shape_name,
                        left,
                        top,
                        width,
                        height,
                    )
                    if hasattr(shape_collection, "_recalculate_extents"):
                        shape_collection._recalculate_extents()
                    shape = shape_collection._shape_factory(textbox)
                if target in {"header", "footer"}:
                    shape.name = f"MCP {target.title()}"
                text_frame = shape.text_frame
                text_frame.word_wrap = True
                text_frame.text = text
                if text_frame.paragraphs and text_frame.paragraphs[0].runs:
                    run = text_frame.paragraphs[0].runs[0]
                    if font_size is not None:
                        run.font.size = Pt(int(font_size))
                    if bold is not None:
                        run.font.bold = bool(bold)
                applied += 1
            continue

        find_text = edit.get("find")
        if not find_text:
            find_text = edit.get("text")
        if not find_text:
            raise ValueError("edit/delete operations require a find or text value")

        replace_text = "" if action == "delete" else str(edit.get("replace", edit.get("text", "")))

        for part in parts:
            shapes = part["shapes"]
            if target == "header":
                tagged = [shape for shape in shapes if shape.name == "MCP Header"]
                _, top, _, height = _pptx_default_box(prs, "header")
                shapes = tagged or [shape for shape in shapes if shape.top < top + height]
            elif target == "footer":
                tagged = [shape for shape in shapes if shape.name == "MCP Footer"]
                shapes = tagged or [shape for shape in shapes if shape.top >= _pptx_default_box(prs, "footer")[1]]
            for text_frame in _pptx_part_text_frames(shapes):
                for paragraph in text_frame.paragraphs:
                    applied += _pptx_replace_in_paragraph(paragraph, re.compile(re.escape(find_text)), replace_text)

    return applied

@mcp.tool()
async def pptx_read(
    path: str,
    structured: bool = False,
    edits: list[dict] | None = None,
    output_path: str | None = None,
) -> list | dict:
    """Read text content from a PowerPoint deck, and optionally apply edits before returning.

    Edits can target slides, notes, slide masters, and layouts, which covers
    normal slide text plus shared header/footer/master-slide content.
    """
    def _run():
        from pptx import Presentation
        source = _resolve(path)
        with _copied_source(source) as source_copy:
            prs = Presentation(str(source_copy))

            if edits:
                applied = _pptx_apply_edits(prs, edits)
                destination = _resolve(output_path) if output_path else source
                destination.parent.mkdir(parents=True, exist_ok=True)
                prs.save(str(destination))
                result = {
                    "updated": str(destination),
                    "edits_applied": applied,
                }
                result["content"] = _pptx_collect_text(prs, structured=structured)
                return result

            return _pptx_collect_text(prs, structured=structured)
    return await asyncio.to_thread(_run)


@mcp.tool()
async def pptx_create(path: str, slides: list[dict], template_path: str | None = None) -> dict:
    """Create a new deck. Optional template_path starts from another PPTX deck.

    Each slide dict: {"title": "...", "body": "..."}.
    """
    def _run():
        import os
        import subprocess
        import tempfile
        from copy import deepcopy
        from io import BytesIO

        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.util import Inches, Pt

        def _add_textbox(slide, left, top, width, height, text, font_size: int | None = None, bold: bool = False):
            shape = slide.shapes.add_textbox(left, top, width, height)
            text_frame = shape.text_frame
            text_frame.word_wrap = True
            text_frame.text = text
            if text_frame.paragraphs and text_frame.paragraphs[0].runs:
                run = text_frame.paragraphs[0].runs[0]
                if font_size is not None:
                    run.font.size = Pt(font_size)
                run.font.bold = bold
            return shape

        def _copy_template_art(source_slide, destination_slide):
            if source_slide is None:
                return

            source_background = source_slide._element.cSld.bg
            if source_background is not None:
                destination_slide._element.cSld.insert(0, deepcopy(source_background))

            for shape in source_slide.shapes:
                if shape.is_placeholder:
                    continue
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    copied_shape = destination_slide.shapes.add_picture(
                        BytesIO(shape.image.blob),
                        shape.left,
                        shape.top,
                        shape.width,
                        shape.height,
                    )
                    try:
                        copied_shape.rotation = shape.rotation
                    except Exception:
                        pass

        def _prepare_template_path(source_path: str) -> tuple[str, str | None]:
            try:
                Presentation(source_path)
                return source_path, None
            except Exception:
                fd, temp_template_path = tempfile.mkstemp(suffix=".pptx")
                os.close(fd)
                powershell_command = (
                    f"Copy-Item -LiteralPath '{source_path}' -Destination '{temp_template_path}' -Force"
                )
                subprocess.run(
                    ["powershell", "-NoProfile", "-Command", powershell_command],
                    check=True,
                    capture_output=True,
                    text=True,
                )
                return temp_template_path, temp_template_path

        p = _resolve(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        cleanup_template_path = None
        if template_path:
            template_source_path, cleanup_template_path = _prepare_template_path(str(_resolve(template_path)))
            template_prs = Presentation(template_source_path)
            source_slide = template_prs.slides[0] if len(template_prs.slides) else None
            prs = Presentation(template_source_path)
            while prs.slides:
                slide_id = prs.slides._sldIdLst[0]
                rel_id = slide_id.rId
                prs.part.drop_rel(rel_id)
                del prs.slides._sldIdLst[0]
        else:
            source_slide = None
            prs = Presentation()

        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        for s in slides:
            slide = prs.slides.add_slide(layout)
            _copy_template_art(source_slide, slide)
            title = s.get("title", "")
            body = s.get("body", "")

            title_shape = slide.shapes.title
            if title_shape is not None:
                title_shape.text = title
            elif title:
                _add_textbox(
                    slide,
                    Inches(0.6),
                    Inches(0.35),
                    prs.slide_width - Inches(1.2),
                    Inches(0.9),
                    title,
                    font_size=28,
                    bold=True,
                )

            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = body
            elif body:
                _add_textbox(
                    slide,
                    Inches(0.8),
                    Inches(1.35),
                    prs.slide_width - Inches(1.6),
                    prs.slide_height - Inches(2.0),
                    body,
                    font_size=20,
                )
        prs.save(str(p))
        result = {"created": str(p), "slides": len(slides)}
        if template_path:
            result["template_used"] = str(_resolve(template_path))
        return result
    
    return await asyncio.to_thread(_run)


@mcp.tool()
async def pptx_modify(
    source_path: str,
    path: str,
    replacements: list[dict] | None = None,
    delete_slide_numbers: list[int] | None = None,
    replacement_slide_range: list[int] | None = None,
    insert_after_slide_number: int | None = None,
    slides: list[dict] | None = None,
    date_all_slides: str | None = None,
    add_text_to_slides: list[dict] | None = None,
    edits: list[dict] | None = None,
    date_as_footer: bool = False,
    create_master: bool = False,
    unhandled_requests: list[str] | None = None,
    modifications: list[dict] | dict | list[str] | str | None = None,
    instruction: str | None = None,
    instructions: str | list[str] | None = None,
    content: str | list | None = None,
) -> dict:
    """Copy an existing PowerPoint deck to a new file and apply modifications.

        - replacements: list of {"find": "...", "replace": "..."} entries.
            Use "find_regex" instead of "find" to match variable text such as amounts.
            Each entry may also include "slide_number" or "slide_numbers" to
            limit that replacement to specific slides.
    - delete_slide_numbers: list of 1-based slide numbers to remove from the copied deck
    - replacement_slide_range: optional [start, end] inclusive 1-based range for replacements
    - insert_after_slide_number: optional 1-based slide number after which new slides are inserted
    - slides: optional list of slide dicts to append to the copied deck
    - date_all_slides: date text to place on every slide, including newly inserted slides
    - add_text_to_slides: list of {"slide_number": 5, "text": "..."} entries, applied after insertion
    - date_as_footer: display date_all_slides as a footer on each slide
    - create_master: add a new slide master via installed Microsoft PowerPoint (Windows only)
        - edits: list of {"action": "add" | "edit" | "delete", "target": "slide" | "header" | "footer" | "master" | "layout", "text": "...", "find": "...", "replace": "..."} entries.
            Optional slide_number, master_index, layout_index and position (in inches) narrow the target.
    """
    def _run():
        import os
        import subprocess
        import tempfile
        from copy import deepcopy
        from io import BytesIO

        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.util import Inches, Pt

        def _add_textbox(slide, left, top, width, height, text, font_size: int | None = None, bold: bool = False, color=None):
            shape = slide.shapes.add_textbox(left, top, width, height)
            text_frame = shape.text_frame
            text_frame.word_wrap = True
            text_frame.text = text
            if text_frame.paragraphs and text_frame.paragraphs[0].runs:
                run = text_frame.paragraphs[0].runs[0]
                if font_size is not None:
                    run.font.size = Pt(font_size)
                run.font.bold = bold
                if color is not None:
                    run.font.color.rgb = color
            return shape

        def _reference_color(slide):
            if slide is not None:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if run.text.strip() and run.font.color.type is not None:
                                    try:
                                        return run.font.color.rgb
                                    except AttributeError:
                                        pass
            return None

        def _copy_background_art(source_slide, destination_slide):
            if source_slide is None:
                return

            source_background = source_slide._element.cSld.bg
            if source_background is not None:
                destination_slide._element.cSld.insert(0, deepcopy(source_background))

            for shape in source_slide.shapes:
                if shape.is_placeholder:
                    continue
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    copied_shape = destination_slide.shapes.add_picture(
                        BytesIO(shape.image.blob),
                        shape.left,
                        shape.top,
                        shape.width,
                        shape.height,
                    )
                    try:
                        copied_shape.rotation = shape.rotation
                    except Exception:
                        pass

        def _prepare_source_copy(source_file_path: str) -> tuple[str, str | None]:
            try:
                Presentation(source_file_path)
                return source_file_path, None
            except Exception:
                fd, temp_source_path = tempfile.mkstemp(suffix=".pptx")
                os.close(fd)
                powershell_command = (
                    f"Copy-Item -LiteralPath '{source_file_path}' -Destination '{temp_source_path}' -Force"
                )
                subprocess.run(
                    ["powershell", "-NoProfile", "-Command", powershell_command],
                    check=True,
                    capture_output=True,
                    text=True,
                )
                return temp_source_path, temp_source_path

        def _normalize_slide_range(slide_range: list[int] | None) -> tuple[int, int] | None:
            if slide_range is None:
                return None
            if len(slide_range) != 2:
                raise ValueError("replacement_slide_range must contain exactly two integers: [start, end]")
            start, end = int(slide_range[0]), int(slide_range[1])
            if start < 1 or end < 1:
                raise ValueError("replacement_slide_range values must be positive slide numbers")
            if start > end:
                raise ValueError("replacement_slide_range start must be less than or equal to end")
            return start, end

        def _slide_in_range(slide_number: int, slide_range: tuple[int, int] | None) -> bool:
            if slide_range is None:
                return True
            start, end = slide_range
            return start <= slide_number <= end

        def _replacement_slides(replacement: dict) -> set[int] | None:
            slide_number = replacement.get("slide_number")
            slide_numbers = replacement.get("slide_numbers")

            if slide_number is not None and slide_numbers is not None:
                raise ValueError("replacement entries may use either slide_number or slide_numbers, not both")

            if slide_number is not None:
                return {int(slide_number)}

            if slide_numbers is not None:
                if isinstance(slide_numbers, int):
                    slide_numbers = [slide_numbers]
                normalized_slide_numbers = {int(number) for number in slide_numbers}
                if not normalized_slide_numbers:
                    raise ValueError("slide_numbers cannot be empty")
                return normalized_slide_numbers

            return None

        def _replacement_applies_to_slide(replacement: dict, slide_number: int) -> bool:
            target_slides = _replacement_slides(replacement)
            return target_slides is None or slide_number in target_slides

        def _insert_slide_after(output_prs, slide, after_slide_number: int):
            if after_slide_number < 1 or after_slide_number > len(output_prs.slides):
                raise ValueError(
                    f"insert_after_slide_number must be between 1 and {len(output_prs.slides)}"
                )

            slide_id = output_prs.slides._sldIdLst[-1]
            output_prs.slides._sldIdLst.remove(slide_id)
            output_prs.slides._sldIdLst.insert(after_slide_number, slide_id)

        def _set_text_on_slide(slide, title: str, body: str, source_prs):
            color = _reference_color(source_slide)
            title_shape = slide.shapes.title
            if title_shape is not None:
                title_shape.text = title
                if color is not None and title_shape.text_frame.paragraphs[0].runs:
                    title_shape.text_frame.paragraphs[0].runs[0].font.color.rgb = color
            elif title:
                _add_textbox(
                    slide,
                    Inches(0.6),
                    Inches(0.35),
                    source_prs.slide_width - Inches(1.2),
                    Inches(0.9),
                    title,
                    font_size=28,
                    bold=True,
                    color=color,
                )

            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = body
                if color is not None and body and slide.placeholders[1].text_frame.paragraphs[0].runs:
                    slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.color.rgb = color
            elif body:
                _add_textbox(
                    slide,
                    Inches(0.8),
                    Inches(1.35),
                    source_prs.slide_width - Inches(1.6),
                    source_prs.slide_height - Inches(2.0),
                    body,
                    font_size=20,
                    color=color,
                )

        # Normalize raw instructions, content, or modifications
        raw_instructions = []
        if isinstance(instruction, str) and instruction.strip():
            raw_instructions.append(instruction.strip())
        if isinstance(instructions, str) and instructions.strip():
            raw_instructions.append(instructions.strip())
        elif isinstance(instructions, list):
            for inst in instructions:
                if isinstance(inst, str) and inst.strip():
                    raw_instructions.append(inst.strip())
        if isinstance(content, str) and content.strip():
            raw_instructions.append(content.strip())

        active_replacements = list(replacements or [])
        active_slides = list(slides or [])
        active_add_text = list(add_text_to_slides or [])
        active_insert_after = insert_after_slide_number
        active_date = date_all_slides
        active_date_as_footer = date_as_footer

        mods_list = []
        if modifications is not None:
            if isinstance(modifications, list):
                mods_list.extend(modifications)
            elif isinstance(modifications, dict):
                mods_list.append(modifications)
            elif isinstance(modifications, str) and modifications.strip():
                raw_instructions.append(modifications.strip())

        for mod in mods_list:
            if isinstance(mod, str):
                raw_instructions.append(mod)
                continue
            if not isinstance(mod, dict):
                continue

            if "find" in mod or "find_regex" in mod or "search" in mod or "old" in mod:
                find_val = mod.get("find") or mod.get("search") or mod.get("old")
                rep_val = mod.get("replace") or mod.get("new") or ""
                entry = {"find": find_val, "replace": rep_val}
                if "slide_number" in mod:
                    entry["slide_number"] = mod["slide_number"]
                if "slide_numbers" in mod:
                    entry["slide_numbers"] = mod["slide_numbers"]
                active_replacements.append(entry)
                continue

            action = str(mod.get("action", "")).lower()
            mod_type = str(mod.get("type", "")).lower()

            if "title" in mod or action in {"add_slide", "insert_slide", "new_slide"} or mod_type == "slide":
                title = mod.get("title", "")
                body = mod.get("body", "")
                active_slides.append({"title": title, "body": body})
                after = mod.get("insert_after") or mod.get("insert_after_slide_number") or mod.get("after") or mod.get("slide_number")
                if after is not None and active_insert_after is None:
                    try:
                        active_insert_after = int(after)
                    except (ValueError, TypeError):
                        pass
                continue

            slide_no = mod.get("slide_number") or mod.get("slide")
            if slide_no is not None and ("text" in mod or "content" in mod or action in {"add_text", "append"}):
                txt = mod.get("text") or mod.get("content") or mod.get("body", "")
                try:
                    active_add_text.append({"slide_number": int(slide_no), "text": str(txt)})
                except (ValueError, TypeError):
                    pass
                continue

            if any(k in mod for k in ("footer", "date", "date_footer")):
                active_date = mod.get("date") or mod.get("footer") or mod.get("date_footer") or "current system date"
                active_date_as_footer = True
                continue

            if "body" in mod or "text" in mod:
                body_content = mod.get("body") or mod.get("text", "")
                rev_m = re.search(r"Revenue\s+Impact:\s*(\$[\d,.]+[a-zA-Z]*)", body_content, re.IGNORECASE)
                if rev_m:
                    active_replacements.append({
                        "find_regex": r"Revenue\s+Impact:\s*\$[\d,.]+[a-zA-Z]*",
                        "replace": f"Revenue Impact: {rev_m.group(1)}",
                    })
                if slide_no is not None:
                    active_add_text.append({"slide_number": int(slide_no), "text": str(body_content)})

        for inst in raw_instructions:
            if re.search(r"\b(?:date|footer)\b", inst, re.IGNORECASE):
                if re.search(r"\b(?:add|insert|include|put)\b.*?\b(?:date|footer)\b", inst, re.IGNORECASE) or "footer" in inst.lower():
                    active_date = "current system date"
                    active_date_as_footer = True

            rev_match = re.search(r"(?:change|update|replace|set)\s+Revenue\s+Impact(?:\s+to|\s+with|\s*=)?\s*(\$[\d,.]+[a-zA-Z]*)", inst, re.IGNORECASE)
            if rev_match:
                active_replacements.append({
                    "find_regex": r"Revenue\s+Impact:\s*\$[\d,.]+[a-zA-Z]*",
                    "replace": f"Revenue Impact: {rev_match.group(1)}",
                })

            for m in re.finditer(r"(?:change|replace)\s+[\"']?([^\"'\n,]+?)[\"']?\s+to\s+[\"']?([^\"'\n,]+?)[\"']?(?:,|$|\.|\s+and)", inst, re.IGNORECASE):
                find_s, repl_s = m.group(1).strip(), m.group(2).strip()
                if "revenue impact" not in find_s.lower():
                    active_replacements.append({"find": find_s, "replace": repl_s})

            slide_match = re.search(r"insert\s+(?:a\s+new\s+)?(?:[“\"']\s*)?([^\"'“”\n]+?)(?:\s*[”\"'])?\s+slide\s+(?:immediately\s+)?after\s+slide\s+(\d+)", inst, re.IGNORECASE)
            if slide_match:
                slide_title = slide_match.group(1).strip()
                after_num = int(slide_match.group(2))
                active_slides.append({"title": slide_title, "body": ""})
                if active_insert_after is None:
                    active_insert_after = after_num

            add_sec_match = re.search(r"add\s+(?:a\s+)?([A-Za-z0-9_\s]+?)\s+section\s+(?:to|in|on)\s+slide\s+(\d+)", inst, re.IGNORECASE)
            if add_sec_match:
                sec_text = add_sec_match.group(1).strip()
                s_num = int(add_sec_match.group(2))
                active_add_text.append({"slide_number": s_num, "text": sec_text})

        source_copy_path = None
        try:
            source_copy_path, source_copy_cleanup = _prepare_source_copy(str(_resolve(source_path)))
            source_prs = Presentation(source_copy_path)
            output_prs = Presentation(source_copy_path)
            source_slide = source_prs.slides[0] if len(source_prs.slides) else None
            normalized_replacement_range = _normalize_slide_range(replacement_slide_range)

            replacement_count = 0
            matched_replacements = set()
            done = []
            not_done = [
                {"request": request, "reason": "Instruction not recognized; use explicit add, update, or delete wording or --args-file"}
                for request in unhandled_requests or []
            ]
            for slide_number, slide in enumerate(output_prs.slides, start=1):
                if not _slide_in_range(slide_number, normalized_replacement_range):
                    continue
                for shape in slide.shapes:
                    if not shape.has_text_frame:
                        continue
                    for paragraph in shape.text_frame.paragraphs:
                        for index, replacement in enumerate(active_replacements):
                            find_text = replacement.get("find", "")
                            find_regex = replacement.get("find_regex", "")
                            replace_text = replacement.get("replace", "")
                            if find_text and not find_regex:
                                if find_text.lower().strip() in {"revenue impact", "revenue impact:"}:
                                    find_regex = r"Revenue\s*Impact:\s*\$[\d,.]+[a-zA-Z]*"
                                    if not replace_text.lower().startswith("revenue"):
                                        replace_text = f"Revenue Impact: {replace_text}"
                            if (find_text or find_regex) and _replacement_applies_to_slide(replacement, slide_number):
                                pattern = re.compile(find_regex if find_regex else re.escape(find_text))
                                count = _pptx_replace_in_paragraph(paragraph, pattern, replace_text)
                                replacement_count += count
                                if count:
                                    matched_replacements.add(index)

            slides_deleted = 0
            if delete_slide_numbers:
                delete_numbers = sorted({int(number) for number in delete_slide_numbers}, reverse=True)
                slide_count = len(output_prs.slides)
                for slide_number in delete_numbers:
                    if slide_number < 1 or slide_number > slide_count:
                        raise ValueError(
                            f"delete_slide_numbers contains out-of-range slide number {slide_number}; deck has {slide_count} slides"
                        )
                for slide_number in delete_numbers:
                    slide_id = output_prs.slides._sldIdLst[slide_number - 1]
                    rel_id = slide_id.rId
                    del output_prs.slides._sldIdLst[slide_number - 1]
                    output_prs.part.drop_rel(rel_id)
                    slide_count -= 1
                    slides_deleted += 1

            slides_added = 0
            if active_slides:
                layout = output_prs.slide_layouts[1] if len(output_prs.slide_layouts) > 1 else output_prs.slide_layouts[0]
                insert_after = active_insert_after
                if insert_after is not None and insert_after < 1:
                    raise ValueError("insert_after_slide_number must be a positive slide number")
                for slide_spec in active_slides:
                    if (insert_after is not None and insert_after < len(output_prs.slides)
                            and any(shape.has_text_frame and shape.text.strip().lower() == slide_spec.get("title", "").lower()
                                    for shape in output_prs.slides[insert_after].shapes)):
                        done.append({"request": f"Insert slide {slide_spec.get('title', '')}", "detail": "Already present after the requested slide"})
                        insert_after += 1
                        continue
                    slide = output_prs.slides.add_slide(layout)
                    _copy_background_art(source_slide, slide)
                    _set_text_on_slide(slide, slide_spec.get("title", ""), slide_spec.get("body", ""), output_prs)
                    if insert_after is not None:
                        _insert_slide_after(output_prs, slide, insert_after)
                        insert_after += 1
                    slides_added += 1
                    done.append({"request": f"Insert slide {slide_spec.get('title', '')}", "detail": "Slide inserted"})

            if len(output_prs.slides) == 0:
                raise ValueError("pptx_modify must leave at least one slide in the deck")

            text_additions_made = 0
            for addition in active_add_text:
                slide_number = int(addition["slide_number"])
                if slide_number < 1 or slide_number > len(output_prs.slides):
                    raise ValueError(f"Slide {slide_number} does not exist in the output deck")
                text = addition["text"].strip()
                if not text:
                    raise ValueError("Slide text to add cannot be empty")
                slide = output_prs.slides[slide_number - 1]
                if any(shape.has_text_frame and text in shape.text for shape in slide.shapes):
                    done.append({"request": f"Add {text} to slide {slide_number}", "detail": "Already present"})
                    continue
                if len(slide.placeholders) > 1 and slide.placeholders[1].has_text_frame:
                    frame = slide.placeholders[1].text_frame
                    if frame.text.strip():
                        paragraph = frame.add_paragraph()
                        paragraph.text = text
                    else:
                        frame.text = text
                        paragraph = frame.paragraphs[0]
                    color = _reference_color(source_slide)
                    if color is not None and paragraph.runs:
                        paragraph.runs[0].font.color.rgb = color
                else:
                    _add_textbox(slide, Inches(0.8), Inches(1.35),
                                 output_prs.slide_width - Inches(1.6), Inches(0.8), text, font_size=20,
                                 color=_reference_color(source_slide))
                text_additions_made += 1
                done.append({"request": f"Add {text} to slide {slide_number}", "detail": "Text added"})

            edits_applied = 0
            for edit in edits or []:
                count = _pptx_apply_edits(output_prs, [edit])
                edits_applied += count
                request = f"{edit.get('action', 'edit')} {edit.get('target', 'slide')}: {edit.get('find', edit.get('text', ''))}"
                if count:
                    done.append({"request": request, "detail": f"Applied {count} time(s)"})
                else:
                    not_done.append({"request": request, "reason": "No matching text or target found"})

            # Auto-resolve current system date if requested
            target_date = active_date
            if target_date or active_date_as_footer:
                import datetime
                date_str = str(target_date).strip() if target_date else ""
                dynamic_tokens = {
                    "current system date", "system date", "current date", "today", "now",
                    "system_date", "current_system_date", "date", "system", "true", "system date footer"
                }
                if not date_str or date_str.lower() in dynamic_tokens or "system date" in date_str.lower() or "current date" in date_str.lower():
                    target_date = datetime.date.today().strftime("%B %d, %Y")
                elif isinstance(target_date, bool) and target_date:
                    target_date = datetime.date.today().strftime("%B %d, %Y")

            if target_date:
                dates_added = 0
                for slide in output_prs.slides:
                    if not any(shape.has_text_frame and target_date in shape.text
                               and (not active_date_as_footer or shape.name == "MCP Footer") for shape in slide.shapes):
                        font_color = _reference_color(slide) or _reference_color(source_slide)
                        date_shape = _add_textbox(
                            slide,
                            output_prs.slide_width - Inches(2.2),
                            output_prs.slide_height - Inches(0.45),
                            Inches(2.0),
                            Inches(0.35),
                            target_date,
                            font_size=10,
                            color=font_color,
                        )
                        date_shape.text_frame.paragraphs[0].alignment = 3
                        if active_date_as_footer:
                            date_shape.name = "MCP Footer"
                        dates_added += 1
                done.append({"request": "Date footer on all slides" if active_date_as_footer else "Date on all slides",
                             "detail": f"Added to {dates_added} slide(s); already present on {len(output_prs.slides) - dates_added}"})

            destination = _resolve(path)
            destination.parent.mkdir(parents=True, exist_ok=True)
            output_prs.save(str(destination))
            done.append({"request": "Save output", "detail": str(destination)})

            for index, replacement in enumerate(active_replacements):
                find_desc = replacement.get("find") or replacement.get("find_regex") or "text"
                request = (f"Replace '{replacement['find']}' with '{replacement.get('replace', '')}'"
                           if replacement.get("find") else f"Update {find_desc} to '{replacement.get('replace', '')}'")
                if index in matched_replacements:
                    done.append({"request": request, "detail": "Text replaced"})
                else:
                    not_done.append({"request": request, "reason": "Text not found on the targeted source slide"})

            if create_master:
                application = presentation = None
                com_initialized = False
                try:
                    import pythoncom
                    from win32com.client import DispatchEx

                    pythoncom.CoInitialize()
                    com_initialized = True
                    with tempfile.TemporaryDirectory() as folder:
                        master_copy = os.path.join(folder, "with-master.pptx")
                        output_prs.save(master_copy)
                        application = DispatchEx("PowerPoint.Application")
                        presentation = application.Presentations.Open(master_copy, ReadOnly=False, Untitled=False, WithWindow=False)
                        presentation.Designs.Add("New Master")
                        presentation.Save()
                        presentation.Close()
                        presentation = None
                        Presentation(master_copy).save(str(destination))
                        if len(Presentation(str(destination)).slide_masters) != len(output_prs.slide_masters) + 1:
                            raise ValueError("New slide master was not present in the saved output")
                        done.append({"request": "Create master slide", "detail": "New slide master created; existing slides remain on their original master"})
                except Exception as exc:
                    not_done.append({"request": "Create master slide", "reason": f"PowerPoint automation unavailable or failed: {exc}"})
                finally:
                    if presentation is not None:
                        presentation.Close()
                    if application is not None:
                        application.Quit()
                    if com_initialized:
                        pythoncom.CoUninitialize()

            return {
                "created": str(destination),
                "source": str(_resolve(source_path)),
                "replacements_made": replacement_count,
                "unmatched_replacements": [
                    replacement for index, replacement in enumerate(active_replacements)
                    if index not in matched_replacements
                ],
                "slides_deleted": slides_deleted,
                "slides_added": slides_added,
                "text_additions_made": text_additions_made,
                "edits_applied": edits_applied,
                "done": done,
                "not_done": not_done,
            }
        finally:
            if source_copy_path and source_copy_path != str(_resolve(source_path)):
                try:
                    os.remove(source_copy_path)
                except OSError:
                    pass

    return await asyncio.to_thread(_run)


@mcp.tool()
async def pptx_add_slide(path: str, title: str, body: str = "") -> dict:
    """Add one slide to an existing deck."""
    def _run():
        from pptx import Presentation
        p = _resolve(path)
        with _copied_source(p) as source_copy:
            prs = Presentation(str(source_copy))
            layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = title
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = body
            prs.save(str(p))
            return {"added_to": str(p), "slide_count": sum(1 for _ in prs.slides)}
    return await asyncio.to_thread(_run)


# --- Smoke test / entry point --------------------------------------------

async def _smoke_test():
    print(f"Registered {len(mcp._tool_manager._tools)} operations:")
    for name in mcp._tool_manager._tools:
        print(f"  - {name}")
    print()
    print(f"FILES_ROOT_DIR: {FILES_ROOT_DIR or '(not set — absolute paths required)'}")
    print("\nTo serve:\n  python files.py --serve\n  python files.py --serve --transport http --port 8008")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--serve", action="store_true")
    parser.add_argument("--transport", choices=["stdio", "http"], default="stdio")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8008)
    args = parser.parse_args()
    if not args.serve:
        asyncio.run(_smoke_test())
        return
    if args.transport == "stdio":
        mcp.run(transport="stdio")
    else:
        mcp.settings.host = args.host
        mcp.settings.port = args.port
        print(f"Serving at http://{args.host}:{args.port}/mcp")
        mcp.run(transport="streamable-http")


if __name__ == "__main__":
    main()
